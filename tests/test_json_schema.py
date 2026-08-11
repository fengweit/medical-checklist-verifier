import json
from importlib.resources import files
from pathlib import Path

from docx import Document
from jsonschema import Draft202012Validator

from medical_checklist_verifier.checklist_digester import digest_checklist


def test_generated_digest_conforms_to_published_json_schema(tmp_path: Path) -> None:
    source = tmp_path / "checklist.docx"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Requirements"
    table.cell(1, 0).text = "1. Identity"
    table.cell(1, 1).text = "Provide manufacturer name."
    document.save(str(source))

    payload = digest_checklist(source)
    schema_path = (
        Path(__file__).parents[1] / "schemas" / "checklist-digester.v1.schema.json"
    )
    schema = json.loads(schema_path.read_text(encoding="utf-8"))

    Draft202012Validator.check_schema(schema)
    Draft202012Validator(schema).validate(payload)


def test_packaged_schema_matches_published_schema() -> None:
    published = (
        Path(__file__).parents[1] / "schemas" / "checklist-digester.v1.schema.json"
    ).read_text(encoding="utf-8")
    packaged = (
        files("medical_checklist_verifier.schemas")
        .joinpath("checklist-digester.v1.schema.json")
        .read_text(encoding="utf-8")
    )

    assert json.loads(packaged) == json.loads(published)
